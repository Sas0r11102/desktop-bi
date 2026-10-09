#!/usr/bin/env python3
"""
Desktop BI 测试数据生成脚本
生成 4 组测试数据，覆盖 CSV + Excel、小数据 + 大数据、多种列类型
"""

import csv
import random
import os
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# 输出目录
OUTPUT_DIR = Path(__file__).parent.parent / "test_data"
OUTPUT_DIR.mkdir(exist_ok=True)

random.seed(42)  # 可复现

# ═══════════════════════════════════════════
# 1. sample_sales.csv — 20 行基础销售数据
# ═══════════════════════════════════════════
def gen_sample_sales():
    filepath = OUTPUT_DIR / "sample_sales.csv"
    categories = ["Electronics", "Clothing", "Food", "Books", "Home"]
    regions = ["North", "South", "East", "West"]
    statuses = ["Completed", "Pending", "Refunded"]

    rows = []
    base_date = datetime(2026, 1, 5)
    for i in range(20):
        date = base_date + timedelta(days=random.randint(0, 90))
        category = random.choice(categories)
        region = random.choice(regions)
        quantity = random.randint(1, 50)
        unit_price = round(random.uniform(10, 500), 2)
        amount = round(quantity * unit_price, 2)
        status = random.choice(statuses)
        rows.append([
            date.strftime("%Y-%m-%d"),
            f"Order-{1001+i}",
            category,
            region,
            quantity,
            unit_price,
            amount,
            status,
        ])

    with open(filepath, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Date", "OrderID", "Category", "Region", "Quantity", "UnitPrice", "Amount", "Status"])
        writer.writerows(rows)

    print(f"✅ {filepath.name} — {len(rows)} rows")


# ═══════════════════════════════════════════
# 2. large_dataset.csv — 1000 行性能测试数据
# ═══════════════════════════════════════════
def gen_large_dataset():
    filepath = OUTPUT_DIR / "large_dataset.csv"
    categories = ["Electronics", "Clothing", "Food", "Books", "Home", "Sports", "Toys", "Beauty"]
    products = [
        "Laptop", "Phone", "Tablet", "Headphones", "Keyboard", "Mouse", "Monitor",
        "T-Shirt", "Jeans", "Jacket", "Sneakers", "Hat", "Socks",
        "Rice", "Milk", "Bread", "Apple", "Banana", "Chicken",
        "Novel", "Textbook", "Magazine", "Comic",
        "Lamp", "Chair", "Desk", "Pillow", "Towel",
        "Basketball", "Tennis Racket", "Yoga Mat",
        "Doll", "Puzzle", "Lego Set",
        "Lipstick", "Perfume", "Sunscreen",
    ]
    regions = ["North", "South", "East", "West", "Central"]
    channels = ["Online", "Offline", "Partner"]
    statuses = ["Completed", "Pending", "Shipped", "Refunded", "Cancelled"]

    rows = []
    base_date = datetime(2026, 1, 1)
    for i in range(1000):
        date = base_date + timedelta(days=i % 365)
        category = random.choice(categories)
        product = random.choice(products)
        region = random.choice(regions)
        channel = random.choice(channels)
        quantity = random.randint(1, 100)
        unit_price = round(random.uniform(5, 2000), 2)
        amount = round(quantity * unit_price, 2)
        discount = round(random.uniform(0, 0.3), 2)
        final_amount = round(amount * (1 - discount), 2)
        status = random.choices(statuses, weights=[60, 15, 10, 10, 5])[0]
        rows.append([
            date.strftime("%Y-%m-%d"),
            f"ORD-{10001+i}",
            category,
            product,
            region,
            channel,
            quantity,
            unit_price,
            amount,
            discount,
            final_amount,
            status,
        ])

    with open(filepath, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow([
            "Date", "OrderID", "Category", "Product", "Region", "Channel",
            "Quantity", "UnitPrice", "Amount", "Discount", "FinalAmount", "Status"
        ])
        writer.writerows(rows)

    size_kb = filepath.stat().st_size / 1024
    print(f"✅ {filepath.name} — {len(rows)} rows ({size_kb:.1f} KB)")


# ═══════════════════════════════════════════
# 3. monthly_finance.xlsx — 多月财务报表（多 Sheet）
# ═══════════════════════════════════════════
def gen_monthly_finance():
    filepath = OUTPUT_DIR / "monthly_finance.xlsx"
    wb = openpyxl.Workbook()

    # 样式
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    num_format = '#,##0.00'

    months = ["January", "February", "March", "April", "May", "June",
              "July", "August", "September", "October", "November", "December"]
    expense_cats = ["Housing", "Food", "Transport", "Shopping", "Entertainment", "Utilities", "Health", "Education"]
    income_sources = ["Salary", "Freelance", "Investment", "Rental", "Bonus"]

    # Sheet 1: Monthly Summary
    ws1 = wb.active
    ws1.title = "Monthly Summary"
    headers1 = ["Month", "Total Income", "Total Expense", "Net Profit", "Savings Rate"]
    ws1.append(headers1)
    for cell in ws1[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for i, month in enumerate(months):
        income = round(random.uniform(8000, 15000), 2)
        expense = round(random.uniform(4000, 9000), 2)
        net = round(income - expense, 2)
        rate = round(net / income * 100, 1) if income > 0 else 0
        ws1.append([month, income, expense, net, rate])
        for cell in ws1[i + 2][1:5]:
            cell.number_format = num_format

    # Sheet 2: Expense Detail
    ws2 = wb.create_sheet("Expense Detail")
    headers2 = ["Month", "Category", "Amount", "Percentage"]
    ws2.append(headers2)
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    row_idx = 2
    for month in months[:6]:  # 只取前 6 个月
        total = round(random.uniform(4000, 9000), 2)
        amounts = []
        for cat in expense_cats:
            amt = round(total * random.uniform(0.05, 0.3), 2)
            amounts.append(amt)
        actual_total = sum(amounts)
        for cat, amt in zip(expense_cats, amounts):
            pct = round(amt / actual_total * 100, 1)
            ws2.append([month, cat, amt, pct])
            ws2.cell(row=row_idx, column=3).number_format = num_format
            row_idx += 1

    # Sheet 3: Income Sources
    ws3 = wb.create_sheet("Income Sources")
    headers3 = ["Month", "Source", "Amount"]
    ws3.append(headers3)
    for cell in ws3[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    row_idx = 2
    for month in months[:6]:
        for src in income_sources:
            amt = round(random.uniform(500, 5000), 2)
            ws3.append([month, src, amt])
            ws3.cell(row=row_idx, column=3).number_format = num_format
            row_idx += 1

    wb.save(filepath)
    size_kb = filepath.stat().st_size / 1024
    print(f"✅ {filepath.name} — 3 sheets ({size_kb:.1f} KB)")


# ═══════════════════════════════════════════
# 4. inventory_items.csv — 库存数据（多种列类型）
# ═══════════════════════════════════════════
def gen_inventory():
    filepath = OUTPUT_DIR / "inventory_items.csv"
    warehouses = ["WH-Shanghai", "WH-Beijing", "WH-Guangzhou", "WH-Chengdu"]
    categories = ["Electronics", "Apparel", "Groceries", "Furniture", "Cosmetics"]
    statuses = ["In Stock", "Low Stock", "Out of Stock", "Discontinued"]
    suppliers = ["Supplier-A", "Supplier-B", "Supplier-C", "Supplier-D", "Supplier-E"]

    rows = []
    for i in range(50):
        sku = f"SKU-{100001+i}"
        name = f"Product-{chr(65 + i % 26)}{i // 26 + 1}"
        category = random.choice(categories)
        warehouse = random.choice(warehouses)
        quantity = random.randint(0, 5000)
        unit_cost = round(random.uniform(5, 800), 2)
        total_value = round(quantity * unit_cost, 2)
        supplier = random.choice(suppliers)
        last_restock = (datetime(2026, 1, 1) + timedelta(days=random.randint(0, 180))).strftime("%Y-%m-%d")
        status = random.choices(statuses, weights=[50, 25, 15, 10])[0]
        is_active = random.choice([True, False])
        rows.append([
            sku, name, category, warehouse, quantity, unit_cost,
            total_value, supplier, last_restock, status, is_active
        ])

    with open(filepath, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow([
            "SKU", "ProductName", "Category", "Warehouse", "Quantity",
            "UnitCost", "TotalValue", "Supplier", "LastRestock", "Status", "IsActive"
        ])
        writer.writerows(rows)

    print(f"✅ {filepath.name} — {len(rows)} rows")


# ═══════════════════════════════════════════
# 执行
# ═══════════════════════════════════════════
if __name__ == "__main__":
    print("📊 生成 Desktop BI 测试数据...\n")
    gen_sample_sales()
    gen_large_dataset()
    gen_monthly_finance()
    gen_inventory()
    print(f"\n📁 输出目录: {OUTPUT_DIR}")
    print("✨ 完成！共 4 个测试数据文件")
